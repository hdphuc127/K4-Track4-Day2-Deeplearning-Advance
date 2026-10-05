"""postprocess.py - bổ sung results.xlsx (độ trễ, chi phí tương đối, Summary nổi bật) và vẽ biểu đồ đánh đổi
độ chính xác - độ trễ. Chạy từ thư mục bài nộp: python code/postprocess.py"""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

X = "results.xlsx"
sh = pd.read_excel(X, sheet_name=None)
lb = pd.read_csv("logs/latency_backbones.csv")
li = pd.read_csv("logs/latency_inference.csv")

# --- Backbones: độ trễ batch 1 (fp32, 224)
f32 = lb[lb.dtype == "fp32"].set_index("config")
fp16 = lb[lb.dtype == "fp16"].set_index("config")
B = sh["Backbones"].copy()
B["lat_p50_ms_b1_fp32"] = B.backbone.map(f32.p50)
B["lat_p95_ms_b1_fp32"] = B.backbone.map(f32.p95)
B["lat_p99_ms_b1_fp32"] = B.backbone.map(f32.p99)
B["lat_p50_ms_b1_fp16"] = B.backbone.map(fp16.p50)

# --- Inference: độ trễ và chi phí so với I00 (chỉ dòng space=prob)
key = {"I00 1-view": "I00 1-view 224", "I01 flip": "I01 flip (K=2) 224", "I02b 5crop+flip192": "I02b 5crop+flip 192 (K=10)",
       "res224": "I00 1-view 224", "res256": "I04 res256", "res288": "I04 res288 (F01)", "res320": "I04 res320",
       "temperature trước (T=0.875)": "I00 1-view 224", "temperature sau (T=0.875)": "I00 1-view 224"}
L = li.set_index("config")
I = sh["Inference"].copy()
I = I[I.space == "prob"].drop(columns="space").reset_index(drop=True)
for c in ("p50", "p95", "p99"):
    I[f"lat_{c}_ms"] = I.method.map(key).map(L[c])
I["cost_vs_I00"] = I.lat_p50_ms / L.loc["I00 1-view 224", "p50"]
I["ghi_chú"] = I.method.map(lambda m: "chưa đo độ trễ (K=5, F1 không hơn I00)" if m == "I02 5crop192" else
                            ("TS: chi phí thêm không đáng kể (một phép chia logit)" if m.startswith("temperature") else ""))
I["K"] = I.method.map({"I01 flip": 2, "I02 5crop192": 5, "I02b 5crop+flip192": 10}).fillna(1).astype(int)

# --- Summary: top theo val + dòng so sánh cuối cùng
S = sh["Summary"].iloc[:, :12].copy()
Fin = sh["Final"]
fin = Fin[Fin.seed.astype(str).str.startswith("mean")].set_index("exp_id")
cmp_ = pd.DataFrame([
    {"cấu hình": "C00 mốc (ConvNeXt-T, 1-view 224)", "macro_f1_test (3 seed)": fin.loc["C00", "macro_f1_test"],
     "top1_test": fin.loc["C00", "top1_test"], "ece_test": fin.loc["C00", "ece_test"], "p95_ms_b1": L.loc["I00 1-view 224", "p95"]},
    {"cấu hình": "F01 tốt nhất (C13 + res288 + TS)", "macro_f1_test (3 seed)": fin.loc["F01", "macro_f1_test"],
     "top1_test": fin.loc["F01", "top1_test"], "ece_test": fin.loc["F01", "ece_test"], "p95_ms_b1": L.loc["I04 res288 (F01)", "p95"]}])

# --- biểu đồ
fig, ax = plt.subplots(1, 2, figsize=(12, 4.6))
a = ax[0]
for _, r in B.iterrows():
    a.scatter(r.lat_p50_ms_b1_fp32, r.val_macro_f1, s=40 + r.params_m * 3, alpha=.8)
    a.annotate(r.backbone, (r.lat_p50_ms_b1_fp32, r.val_macro_f1), fontsize=8, xytext=(4, 4), textcoords="offset points")
a.set(xlabel="p50 batch 1, FP32, 224 (ms, T4)", ylabel="macro-F1 val", title="Backbone (diện tích ∝ số tham số)")
a.grid(alpha=.3)
b = ax[1]
P = I.dropna(subset=["lat_p50_ms"]).drop_duplicates("method")
P = P[~P.method.str.startswith("temperature")]
b.scatter(P.lat_p50_ms, P.macro_f1_val, c="tab:red")
for _, r in P.iterrows():
    b.annotate(r.method, (r.lat_p50_ms, r.macro_f1_val), fontsize=8, xytext=(4, -9), textcoords="offset points")
b.set(xlabel="p50 (ms, T4, FP32, theo K view/độ phân giải)", ylabel="macro-F1 val", title="Suy luận trên ConvNeXt-T (C13)")
b.grid(alpha=.3)
fig.tight_layout()
fig.savefig("figures/tradeoff_acc_latency.png", dpi=140)

# --- ghi lại
sheets = {**sh, "Backbones": B, "Inference": I, "Summary": S, "Latency_extra": pd.concat(
    [lb.assign(nhóm="backbone b1"), li.assign(nhóm="inference ConvNeXt-T")]), "Final_vs_baseline": cmp_}
with pd.ExcelWriter(X, engine="openpyxl") as w:
    for n, d in sheets.items():
        d.round(4).to_excel(w, sheet_name=n, index=False)
wb = load_workbook(X)
green = PatternFill("solid", fgColor="C6EFCE")
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for c in ws[1]:
        c.font = Font(bold=True)
ws = wb["Summary"]            # dòng tốt nhất theo val + mốc C00
for row in ws.iter_rows(min_row=2):
    if row[0].value in ("C13", "C00"):
        for c in row: c.fill = green if row[0].value == "C13" else PatternFill("solid", fgColor="FFEB9C")
wsf = wb["Final_vs_baseline"]
for c in wsf[3]: c.fill = green
wb.save(X)
print("ok", B[["backbone", "lat_p50_ms_b1_fp32", "lat_p95_ms_b1_fp32"]].round(2).to_string())
